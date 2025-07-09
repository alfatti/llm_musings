#!/usr/bin/env python3
"""
End‑to‑end rent‑roll extraction pipeline using **Gemini Pro**.

Workflow
~~~~~~~~
1. **Load Excel ➜ dense grid** (no schema assumptions).
2. **Phase A – Header discovery**: ask Gemini to find the canonical header row and map each header to a set of canonical keys (property, unit, ...).
3. **Phase B – Row extraction**: stream grid chunks (≤ CHUNK_SIZE rows) back to Gemini together with the mapping; receive JSON records.
4. **Validator / Retry**: enforce a strict schema locally; auto‑retry bad chunks up to MAX_RETRIES.
5. **Assemble & save**: concatenates records → pandas DataFrame, writes `*.json` + `*.parquet`.

The prompts come from our previous chats and are embedded below as Jinja‑style templates.
"""
from __future__ import annotations

import itertools
import json
import os
import sys
from pathlib import Path
from typing import Any, Dict, List

import openpyxl  # pip install openpyxl
import pandas as pd  # pip install pandas
from jsonschema import Draft7Validator, ValidationError  # pip install jsonschema
from langchain_google_genai import ChatGoogleGenerativeAI  # pip install langchain-google-genai

# ────────────────────────────────────────────
# Configuration
# ────────────────────────────────────────────
CHUNK_SIZE: int = 200  # rows per LLM call during Phase B
MAX_RETRIES: int = 3
MODEL_NAME: str = "gemini-1.5-pro-latest"  # or gemini-2.5-flash‑preview‑04‑17

# Canonical keys we expect for each tenant record
CANONICAL_KEYS: List[str] = [
    "property",
    "unit",
    "market_rent",
    "current_rent",
    "lease_start",
    "lease_end",
    "gla",
]

# JSONSchema for validator / retry loop
SCHEMA: Dict[str, Any] = {
    "type": "array",
    "items": {
        "type": "object",
        "properties": {
            "property": {"type": "string"},
            "unit": {"type": ["string", "number"]},
            "market_rent": {"type": ["string", "number"]},
            "current_rent": {"type": ["string", "number"]},
            "lease_start": {"type": "string"},
            "lease_end": {"type": "string"},
            "gla": {"type": ["string", "number"]},
        },
        "required": CANONICAL_KEYS,
        "additionalProperties": False,
    },
}

validator = Draft7Validator(SCHEMA)

# ────────────────────────────────────────────
# Prompt templates (Jinja‑style placeholders)
# ────────────────────────────────────────────
HEADER_PROMPT_TEMPLATE = """<system>
You are a data‑extraction expert.
</system>
<user>
Below is the first {{n_rows}} rows of an Excel sheet serialized as JSON under the key `cells`.

```json
{"cells": {{rows_json}}}
```

1 – Identify the **canonical header row** (first row that looks like column names).
2 – Map each header you find to one of the canonical keys:
   property, unit, market_rent, current_rent, lease_start, lease_end, gla
   (ignore columns that don't match any key).
3 – Return **exactly** the JSON object:
```json
{ "header_row": <int>, "mapping": { "<canonical_key>": "<detected header text>", … } }
```
</user>"""

ROW_EXTRACT_PROMPT_TEMPLATE = """<user>
Extract tenant rows.

**Header row index**: {{header_row}}
**Column map**: {{column_map}}

Return **only** a JSON list where each element has keys
property, unit, market_rent, current_rent, lease_start, lease_end, gla.
Use string values unless a numeric type is obvious.
</user>

```json
{{rows_json}}
```"""

# ────────────────────────────────────────────
# Gemini client helper
# ────────────────────────────────────────────

def make_gemini(model_name: str = MODEL_NAME) -> ChatGoogleGenerativeAI:
    """Instantiate a LangChain wrapper around Gemini Pro.

    The following environment variables must be set *or* managed via Vault:
        GOOGLE_API_KEY, GOOGLE_CLOUD_LOCATION, GOOGLE_CLOUD_PROJECT, ...
    """
    return ChatGoogleGenerativeAI(model=model_name)


def call_llm(llm: ChatGoogleGenerativeAI, prompt: str) -> Any:
    """Send prompt, return parsed JSON."""
    raw = llm.invoke(prompt)
    # LangChain returns a BaseMessage; `.content` holds the text.
    txt = raw.content if hasattr(raw, "content") else raw
    try:
        return json.loads(txt)
    except json.JSONDecodeError as exc:
        raise RuntimeError(f"LLM response is not valid JSON: {txt[:200]}…") from exc

# ────────────────────────────────────────────
# Excel utilities
# ────────────────────────────────────────────

def excel_to_grid(path: Path) -> List[List[str]]:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    grid: List[List[str]] = []
    max_nonempty = -1
    for row in ws.iter_rows():
        r = [str(c.value).strip() if c.value is not None else "" for c in row]
        # track last non‑empty column across the sheet (token saver)
        if any(r):
            max_nonempty = max(max_nonempty, max(i for i, v in enumerate(r) if v))
        grid.append(r)
    if max_nonempty >= 0:
        grid = [r[: max_nonempty + 1] for r in grid]
    return grid

# ────────────────────────────────────────────
# Phase A – Header detection
# ────────────────────────────────────────────

def detect_header(llm: ChatGoogleGenerativeAI, grid: List[List[str]], sample_rows: int = 40) -> tuple[int, Dict[str, str]]:
    excerpt = grid[:sample_rows]
    prompt = HEADER_PROMPT_TEMPLATE.replace("{{n_rows}}", str(sample_rows)).replace(
        "{{rows_json}}", json.dumps(excerpt, separators=",:")
    )
    resp = call_llm(llm, prompt)
    header_row = int(resp["header_row"])  # type: ignore[index]
    mapping = resp["mapping"]  # type: ignore[index]
    return header_row, mapping


def build_index_map(mapping: Dict[str, str], header_row: List[str]) -> Dict[str, int]:
    """Convert header‑text map → column index map."""
    idx_map: Dict[str, int] = {}
    for canonical, header_text in mapping.items():
        try:
            idx_map[canonical] = next(i for i, v in enumerate(header_row) if v == header_text)
        except StopIteration:
            # fallback: best‑effort case‑insensitive match
            idx_map[canonical] = next(
                i for i, v in enumerate(header_row) if v.lower() == header_text.lower()
            )
    return idx_map

# ────────────────────────────────────────────
# Phase B – Row extraction, validation & retry
# ────────────────────────────────────────────

def extract_chunk(
    llm: ChatGoogleGenerativeAI,
    header_row_num: int,
    idx_map: Dict[str, int],
    rows: List[List[str]],
) -> List[Dict[str, Any]]:
    tmpl = ROW_EXTRACT_PROMPT_TEMPLATE
    prompt = (
        tmpl.replace("{{header_row}}", str(header_row_num))
        .replace("{{column_map}}", json.dumps(idx_map))
        .replace("{{rows_json}}", json.dumps(rows, separators=",:"))
    )
    return call_llm(llm, prompt)


def validate_records(records: List[Dict[str, Any]]) -> None:
    errors = sorted(validator.iter_errors(records), key=lambda e: e.path)
    if errors:
        raise ValidationError("; ".join(e.message for e in errors))

# ────────────────────────────────────────────
# Main driver
# ────────────────────────────────────────────

def process_excel(path: Path, out_dir: Path | None = None) -> Path:
    out_dir = out_dir or path.parent
    llm = make_gemini()

    grid = excel_to_grid(path)
    header_row_num, mapping = detect_header(llm, grid)
    header_row = grid[header_row_num]
    idx_map = build_index_map(mapping, header_row)

    # Stream after header row
    data_rows = grid[header_row_num + 1 :]
    records: List[Dict[str, Any]] = []

    for chunk_idx, i in enumerate(range(0, len(data_rows), CHUNK_SIZE)):
        chunk = data_rows[i : i + CHUNK_SIZE]
        attempt = 0
        while attempt < MAX_RETRIES:
            attempt += 1
            try:
                recs = extract_chunk(llm, header_row_num, idx_map, chunk)
                validate_records(recs)
                records.extend(recs)
                break  # success → next chunk
            except (ValidationError, RuntimeError) as exc:
                print(f"Chunk {chunk_idx} attempt {attempt} failed: {exc}")
                if attempt >= MAX_RETRIES:
                    print("❌ giving up on this chunk; logging for review …")
                    Path(out_dir / f"failed_chunk_{chunk_idx}.json").write_text(
                        json.dumps({"rows": chunk}, indent=2)
                    )

    # Assemble outputs
    out_json = out_dir / f"{path.stem}_records.json"
    out_parquet = out_dir / f"{path.stem}_records.parquet"

    with out_json.open("w") as fp:
        json.dump(records, fp, indent=2)

    df = pd.DataFrame(records)
    df.to_parquet(out_parquet, index=False)
    print(f"✅ Saved {len(records)} records → {out_json} & {out_parquet}")
    return out_json


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description="Extract rent‑roll data via Gemini Pro")
    parser.add_argument("excel_file", type=Path, help="Path to the rent‑roll .xlsx file")
    parser.add_argument("--out", type=Path, default=None, help="Output directory (default: same as input)")
    parser.add_argument("--chunk", type=int, default=CHUNK_SIZE, help="Rows per LLM call")
    parser.add_argument("--retries", type=int, default=MAX_RETRIES, help="Max retries per chunk")
    args = parser.parse_args()

    CHUNK_SIZE = args.chunk
    MAX_RETRIES = args.retries

    try:
        process_excel(args.excel_file, args.out)
    except Exception as e:
        print(f"Fatal error: {e}")
        sys.exit(1)
